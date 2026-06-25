from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path

from ..config import get_settings
from ..db import get_session
from ..models import Document
from .store import add_chunks, delete_document

_settings = get_settings()


def _hash_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _chunk(text: str, size: int, overlap: int) -> list[str]:
    text = " ".join(text.split())
    out, start = [], 0
    while start < len(text):
        out.append(text[start : start + size])
        start = max(start + size - overlap, start + 1)
    return [c for c in out if c.strip()]


def extract_text(path: Path) -> list[tuple[int, str]]:
    """Return [(page_number, text), ...] for supported file types."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            return [(i + 1, page.get_text()) for i, page in enumerate(doc)]
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        return [(1, "\n".join(p.text for p in d.paragraphs))]
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for idx, ws in enumerate(wb.worksheets, start=1):
            rows = [
                " ".join(str(c) for c in row if c is not None)
                for row in ws.iter_rows(values_only=True)
            ]
            out.append((idx, "\n".join(rows)))
        return out
    if suffix in (".txt", ".md"):
        return [(1, path.read_text(errors="ignore"))]
    raise ValueError(f"Unsupported file type: {suffix}")


def ingest_document(doc_id: int) -> dict:
    with get_session() as session:
        doc = session.get(Document, doc_id)
        if not doc:
            return {"ok": False, "error": "document not found"}
        path = _settings.documents_dir / doc.path
        if not path.exists():
            return {"ok": False, "error": f"file missing: {path}"}

        delete_document(doc_id)  # idempotent re-ingest

        all_chunks: list[str] = []
        all_pages: list[int] = []
        for page_no, text in extract_text(path):
            for ch in _chunk(text, _settings.chunk_size, _settings.chunk_overlap):
                all_chunks.append(ch)
                all_pages.append(page_no)

        add_chunks(doc_id, doc.machine_ids, doc.title, doc.doc_type, all_chunks, all_pages)

        doc.content_hash = _hash_file(path)
        doc.chunk_count = len(all_chunks)
        doc.ingested_at = datetime.now(timezone.utc)
        session.add(doc)
        session.commit()
        return {"ok": True, "chunks": len(all_chunks)}
